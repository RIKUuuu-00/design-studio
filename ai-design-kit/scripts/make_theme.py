#!/usr/bin/env python3
"""デザイン方針（theme.json）から、社内テンプレが無いとき用のテンプレ（.pptx）を作る。

    python make_theme.py design-system/slides/theme.json --out design-system/slides/template.pptx

作ったテンプレは社内テンプレと同じく register_template.py → build_deck.py で使う。
テーマ色・テーマフォント・マスターの文字スタイル・レイアウト（表紙 / 章扉 / 本文 / 2列 / タイトルのみ / 白紙）を
持つ本物のテンプレなので、PowerPoint で開いてもテーマとして編集できる。

theme.json:
{
  "name": "Harbor Navy",
  "colors": {                       # テーマの 12 色（# なし）。dk1=本文文字, lt1=背景, dk2=濃い基調色, lt2=淡い基調色
    "dk1": "1F2328", "lt1": "FFFFFF", "dk2": "0E2F4F", "lt2": "EEF1F4",
    "accent1": "1F5FAD", "accent2": "D9822B", "accent3": "5B8C5A",
    "accent4": "8A6FB0", "accent5": "3A9BA8", "accent6": "B5524A",
    "hlink": "1F5FAD", "folHlink": "6B4E99"
  },
  "fonts": {"heading": {"latin": "Noto Sans JP", "ea": "Noto Sans JP"},
            "body":    {"latin": "Noto Sans JP", "ea": "Noto Sans JP"}},
  "title_size": 26, "body_size": 18,
  "title_color": "dk1",             # 本文スライドのタイトル色（dk1 / dk2 / accent1）
  "cover": "split",                 # 表紙: split（左に色面） / full（全面に色） / minimal（白地に大きな文字）
  "section": "full",                # 章扉: full / minimal
  "title_rule": true,               # 本文スライドのタイトル下に細い区切り線
  "size": "16:9"                    # 16:9 / 4:3
}
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib.pptx_common import NS  # noqa: E402

SLOTS = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"]
SCHEME_REF = {"dk1": "tx1", "lt1": "bg1", "dk2": "tx2", "lt2": "bg2"}


# ---------------------------------------------------------------- 色の検査

def _lum(hex6: str) -> float:
    def ch(v):
        v = v / 255
        return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (int(hex6[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * ch(r) + 0.7152 * ch(g) + 0.0722 * ch(b)


def contrast(a: str, b: str) -> float:
    la, lb = sorted((_lum(a), _lum(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def check_palette(colors: dict) -> list[str]:
    warn = []
    rules = [
        ("dk1", "lt1", 7.0, "本文文字と背景"),
        ("dk2", "lt1", 4.5, "濃い基調色と背景"),
        ("lt1", "dk2", 4.5, "章扉・表紙の白文字と濃い基調色"),
        ("accent1", "lt1", 3.0, "強調色と背景（大きな文字・図形）"),
        ("lt1", "accent1", 3.0, "強調色の面に載せる白文字"),
    ]
    for fg, bg, need, label in rules:
        c = contrast(colors[fg], colors[bg])
        if c < need:
            warn.append(f"{label}のコントラスト {c:.1f}:1（目安 {need}:1 以上）。{fg} か {bg} を調整する")
    # グラフで使う色同士が近すぎないか（知覚的な色差 ΔE。20 未満は見分けにくい）
    acc = [colors[f"accent{i}"] for i in range(1, 7)]
    for i in range(len(acc)):
        for j in range(i + 1, len(acc)):
            d = delta_e(acc[i], acc[j])
            if d < 20:
                warn.append(f"accent{i + 1} と accent{j + 1} の色差が小さい（ΔE {d:.0f}）。グラフで見分けにくい")
    return warn


def _lab(hex6: str) -> tuple[float, float, float]:
    def lin(v):
        v = v / 255
        return v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4
    r, g, b = (lin(int(hex6[i:i + 2], 16)) for i in (0, 2, 4))
    x = (0.4124 * r + 0.3576 * g + 0.1805 * b) / 0.95047
    y = 0.2126 * r + 0.7152 * g + 0.0722 * b
    z = (0.0193 * r + 0.1192 * g + 0.9505 * b) / 1.08883
    def f(t):
        return t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116
    fx, fy, fz = f(x), f(y), f(z)
    return 116 * fy - 16, 500 * (fx - fy), 200 * (fy - fz)


def delta_e(a: str, b: str) -> float:
    la, lb = _lab(a), _lab(b)
    return sum((p - q) ** 2 for p, q in zip(la, lb)) ** 0.5


# ---------------------------------------------------------------- テーマ

def set_theme(prs, plan: dict) -> None:
    master = prs.slide_masters[0]
    part = next(r.target_part for r in master.part.rels.values() if r.reltype.endswith("/theme"))
    root = etree.fromstring(part.blob)
    root.set("name", plan.get("name", "Custom"))
    scheme = root.find(".//a:clrScheme", NS)
    scheme.set("name", plan.get("name", "Custom"))
    colors = plan["colors"]
    for slot in SLOTS:
        el = scheme.find(f"a:{slot}", NS)
        for c in list(el):
            el.remove(c)
        etree.SubElement(el, qn("a:srgbClr"), val=colors[slot].upper().lstrip("#"))
    fonts = plan["fonts"]
    fs = root.find(".//a:fontScheme", NS)
    fs.set("name", plan.get("name", "Custom"))
    for kind, key in (("major", "heading"), ("minor", "body")):
        f = fs.find(f"a:{kind}Font", NS)
        f.find("a:latin", NS).set("typeface", fonts[key]["latin"])
        f.find("a:ea", NS).set("typeface", fonts[key]["ea"])
        # スクリプト別の上書き（Jpan など）は消す。テーマの ea を常に使わせるため
        for extra in f.findall("a:font", NS):
            f.remove(extra)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ---------------------------------------------------------------- 文字スタイル

def _solid(parent, ref: str):
    fill = etree.SubElement(parent, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:schemeClr"), val=SCHEME_REF.get(ref, ref))
    return fill


def _set_lvl(lvl, size: float | None = None, bold: bool | None = None, color: str | None = None,
             align: str | None = None, bullet: str | None = None, mar: int | None = None, before: float | None = None,
             line: float | None = None, heading: bool | None = None):
    """lvlNpPr を作り直す（子要素の順序: lnSpc, spcBef, spcAft, buClr, buSzPct, buFont, bu*, defRPr）。"""
    for c in list(lvl):
        lvl.remove(c)
    if align:
        lvl.set("algn", align)
    if mar is not None:
        lvl.set("marL", str(mar))
        lvl.set("indent", str(-mar if bullet else 0))
    if line:
        ln = etree.SubElement(lvl, qn("a:lnSpc"))
        etree.SubElement(ln, qn("a:spcPct"), val=str(int(line * 100000)))
    if before is not None:
        sb = etree.SubElement(lvl, qn("a:spcBef"))
        etree.SubElement(sb, qn("a:spcPts"), val=str(int(before * 100)))
    if bullet:
        bc = etree.SubElement(lvl, qn("a:buClr"))
        etree.SubElement(bc, qn("a:schemeClr"), val="accent1")
        etree.SubElement(lvl, qn("a:buFont"), typeface="Arial")
        etree.SubElement(lvl, qn("a:buChar"), char=bullet)
    elif bullet == "":
        etree.SubElement(lvl, qn("a:buNone"))
    rpr = etree.SubElement(lvl, qn("a:defRPr"))
    if size:
        rpr.set("sz", str(int(size * 100)))
    if bold is not None:
        rpr.set("b", "1" if bold else "0")
    if color:
        _solid(rpr, color)
    face = "+mj" if heading else "+mn"
    etree.SubElement(rpr, qn("a:latin"), typeface=f"{face}-lt")
    etree.SubElement(rpr, qn("a:ea"), typeface=f"{face}-ea")
    etree.SubElement(rpr, qn("a:cs"), typeface=f"{face}-cs")


def style_master(master, plan: dict) -> None:
    tx = master._element.find("p:txStyles", NS)
    title = tx.find("p:titleStyle", NS)
    _set_lvl(title.find("a:lvl1pPr", NS), size=plan.get("title_size", 26), bold=True, color=plan.get("title_color", "dk1"),
             align="l", line=1.1, heading=True)
    body = tx.find("p:bodyStyle", NS)
    base = plan.get("body_size", 18)
    sizes = [base, base - 2, base - 4, base - 4, base - 4]
    for i in range(1, 10):
        lvl = body.find(f"a:lvl{i}pPr", NS)
        if lvl is None:
            continue
        s = sizes[min(i - 1, 4)]
        _set_lvl(lvl, size=s, color="dk1", align="l", bullet="•" if i == 1 else "–", mar=int(Pt(s) * (1.1 + 1.3 * (i - 1))),
                 before=s * 0.45, line=1.05)


def _ph(container, ptype):
    return next((p for p in container.placeholders if p.placeholder_format.type == ptype), None)


def _place(shape, x, y, w, h):
    shape.left, shape.top, shape.width, shape.height = Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h))


def _anchor(shape, anchor: str):
    body_pr = shape._element.find(".//a:bodyPr", NS)
    if body_pr is not None:
        body_pr.set("anchor", anchor)


def _ph_style(shape, size=None, bold=None, color=None, align=None, heading=None, bullet=None):
    """プレースホルダー自身の lstStyle で文字を上書きする（レイアウト単位の調整）。"""
    tx_body = shape._element.find("p:txBody", NS)
    lst = tx_body.find("a:lstStyle", NS)
    if lst is None:
        lst = etree.Element(qn("a:lstStyle"))
        tx_body.insert(1, lst)
    for c in list(lst):
        lst.remove(c)
    lvl = etree.SubElement(lst, qn("a:lvl1pPr"))
    _set_lvl(lvl, size=size, bold=bold, color=color, align=align, heading=heading, bullet=bullet, mar=0 if bullet == "" else None)


def _layout_shape(layout, tmp_slide, kind: str, box, color: str, width_pt: float = 0.75):
    """レイアウトに装飾図形を置く（LayoutShapes は追加 API を持たないため、仮スライドで作って移す）。"""
    x, y, w, h = (int(v) for v in box)
    if kind == "line":
        shp = tmp_slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x), Emu(y), Emu(x + w), Emu(y + h))
        shp.line.width = Pt(width_pt)
        shp.line.color.theme_color = _theme_enum(color)
    else:
        shp = tmp_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
        shp.fill.solid()
        shp.fill.fore_color.theme_color = _theme_enum(color)
        shp.line.fill.background()
        shp.shadow.inherit = False
    el = shp._element
    style = el.find(qn("p:style"))
    if style is not None:
        el.remove(style)
    el.getparent().remove(el)
    ids = [int(v) for v in layout._element.xpath("//p:cNvPr/@id")]
    el.xpath("./*[1]/p:cNvPr")[0].set("id", str(max(ids + [1]) + 1))
    # 装飾はプレースホルダーより背面に、追加した順に重ねる（最初のプレースホルダーの直前に入れる）
    tree = layout.shapes._spTree
    first_ph = next((c for c in tree if c.find(".//p:nvPr/p:ph", NS) is not None), None)
    if first_ph is not None:
        first_ph.addprevious(el)
    else:
        tree.append(el)


def _theme_enum(name: str):
    from pptx.enum.dml import MSO_THEME_COLOR

    return {
        "dk1": MSO_THEME_COLOR.TEXT_1, "lt1": MSO_THEME_COLOR.BACKGROUND_1, "dk2": MSO_THEME_COLOR.TEXT_2,
        "lt2": MSO_THEME_COLOR.BACKGROUND_2, **{f"accent{i}": getattr(MSO_THEME_COLOR, f"ACCENT_{i}") for i in range(1, 7)},
    }[name]


def _bg(layout, color: str) -> None:
    fill = layout.background.fill
    fill.solid()
    fill.fore_color.theme_color = _theme_enum(color)


# ---------------------------------------------------------------- 組み立て

def build(plan: dict):
    prs = Presentation()
    if plan.get("size", "16:9") == "16:9":
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = int(prs.slide_width), int(prs.slide_height)
    M = int(W * 0.045)                       # 左右の余白
    title_y, title_h = int(H * 0.06), int(H * 0.15)
    body_y = int(H * 0.25)
    foot_y, foot_h = int(H * 0.92), int(H * 0.05)
    body_h = foot_y - int(H * 0.03) - body_y
    set_theme(prs, plan)
    colors = {k: v.lstrip("#") for k, v in plan["colors"].items()}
    # 濃い面に置く目印の色: accent1 が沈むなら accent2、それも沈むなら白
    mark = next((c for c in ("accent1", "accent2") if contrast(colors[c], colors["dk2"]) >= 2.2), "lt1")
    inset = int(0.1 * 914400)  # テキスト枠の内部余白。目印を文字の左端にそろえる
    master = prs.slide_masters[0]
    style_master(master, plan)

    # マスターの枠位置
    _place(_ph(master, PP_PLACEHOLDER.TITLE), M, title_y, W - 2 * M, title_h)
    _anchor(_ph(master, PP_PLACEHOLDER.TITLE), "b")
    _place(_ph(master, PP_PLACEHOLDER.BODY), M, body_y, W - 2 * M, body_h)
    foot_size = max(plan.get("body_size", 18) - 8, 9)
    for ptype, x, w, align in (
        (PP_PLACEHOLDER.FOOTER, M, W * 0.5, "l"),
        (PP_PLACEHOLDER.DATE, W * 0.55, W * 0.2, "r"),
        (PP_PLACEHOLDER.SLIDE_NUMBER, W - M - W * 0.08, W * 0.08, "r"),
    ):
        ph = _ph(master, ptype)
        _place(ph, x, foot_y, w, foot_h)
        _ph_style(ph, size=foot_size, color="dk1", align=align)
        lvl_rpr = ph._element.find(".//a:lstStyle/a:lvl1pPr/a:defRPr/a:solidFill/a:schemeClr", NS)
        etree.SubElement(lvl_rpr, qn("a:lumMod"), val="50000")
        etree.SubElement(lvl_rpr, qn("a:lumOff"), val="50000")

    # 不要なレイアウトを外す（縦書き・キャプション付き・比較）
    keep = {"Title Slide": "表紙", "Title and Content": "本文", "Section Header": "章扉", "Two Content": "2列",
            "Title Only": "タイトルのみ", "Blank": "白紙"}
    for layout in list(prs.slide_layouts):
        if layout.name not in keep:
            prs.slide_layouts.remove(layout)
    tmp = prs.slides.add_slide(prs.slide_layouts[0])  # 装飾図形の作成用（最後に消す）

    for layout in prs.slide_layouts:
        name = layout.name
        layout.name = keep[name]
        title = _ph(layout, PP_PLACEHOLDER.TITLE) or _ph(layout, PP_PLACEHOLDER.CENTER_TITLE)
        if name in ("Title and Content", "Two Content", "Title Only"):
            _place(title, M, title_y, W - 2 * M, title_h)
            _anchor(title, "b")
            if plan.get("title_rule", True):
                _layout_shape(layout, tmp, "line", (M, title_y + title_h + int(H * 0.025), W - 2 * M, 0), "dk1", 0.5)
        if name == "Title and Content":
            body = [p for p in layout.placeholders if p.placeholder_format.idx == 1][0]
            _place(body, M, body_y, W - 2 * M, body_h)
        if name == "Two Content":
            gap = int(W * 0.035)
            bw = (W - 2 * M - gap) // 2
            bodies = sorted([p for p in layout.placeholders if p.placeholder_format.type in (PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.BODY)],
                            key=lambda p: p.left)
            for i, b in enumerate(bodies[:2]):
                _place(b, M + i * (bw + gap), body_y, bw, body_h)
        if name == "Title Slide":
            sub = _ph(layout, PP_PLACEHOLDER.SUBTITLE)
            cover = plan.get("cover", "split")
            big = plan.get("title_size", 26) + 10
            if cover == "split":
                panel_w = int(W * 0.6)
                _layout_shape(layout, tmp, "rect", (0, 0, panel_w, H), "dk2")
                _layout_shape(layout, tmp, "rect", (M + inset, int(H * 0.30), int(W * 0.05), int(H * 0.012)), mark)
                _place(title, M, int(H * 0.34), panel_w - 2 * M, int(H * 0.34))
                _place(sub, M, int(H * 0.72), panel_w - 2 * M, int(H * 0.12))
                fg, sub_color = "lt1", "lt2"
            elif cover == "full":
                _bg(layout, "dk2")
                _layout_shape(layout, tmp, "rect", (M + inset, int(H * 0.30), int(W * 0.05), int(H * 0.012)), mark)
                _place(title, M, int(H * 0.34), int(W * 0.72), int(H * 0.32))
                _place(sub, M, int(H * 0.70), int(W * 0.6), int(H * 0.12))
                fg, sub_color = "lt1", "lt2"
            else:  # minimal
                _layout_shape(layout, tmp, "rect", (M + inset, int(H * 0.30), int(W * 0.05), int(H * 0.012)), "accent1")
                _place(title, M, int(H * 0.34), int(W * 0.8), int(H * 0.32))
                _place(sub, M, int(H * 0.70), int(W * 0.6), int(H * 0.12))
                fg, sub_color = "dk1", "dk2"
            _anchor(title, "t")
            _anchor(sub, "t")
            _ph_style(title, size=big, bold=True, color=fg, align="l", heading=True)
            _ph_style(sub, size=plan.get("body_size", 18), bold=False, color=sub_color, align="l", bullet="")
            for p in list(layout.placeholders):
                if p.placeholder_format.type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER):
                    p._element.getparent().remove(p._element)
        if name == "Section Header":
            label = _ph(layout, PP_PLACEHOLDER.BODY)
            dark = plan.get("section", "full") == "full"
            if dark:
                _bg(layout, "dk2")
            _layout_shape(layout, tmp, "rect", (M + inset, int(H * 0.36), int(W * 0.05), int(H * 0.012)), mark if dark else "accent1")
            _place(label, M, int(H * 0.26), int(W * 0.6), int(H * 0.08))
            _place(title, M, int(H * 0.40), int(W * 0.8), int(H * 0.30))
            _anchor(title, "t")
            _anchor(label, "b")
            _ph_style(title, size=plan.get("title_size", 26) + 8, bold=True, color="lt1" if dark else "dk1", align="l", heading=True)
            _ph_style(label, size=plan.get("body_size", 18) - 2, bold=True, color="lt2" if dark else "accent1", align="l", bullet="")
            if dark:
                for p in layout.placeholders:
                    if p.placeholder_format.type in (PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER):
                        rpr = p._element.find(".//a:lstStyle/a:lvl1pPr/a:defRPr", NS)
                        if rpr is None:
                            _ph_style(p, size=foot_size, color="lt2", align="r" if p.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER else "l")
    # 仮スライドを外す
    sld_id = prs.slides._sldIdLst[0]
    prs.part.drop_rel(sld_id.rId)
    prs.slides._sldIdLst.remove(sld_id)
    return prs


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("theme_json")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    plan = json.loads(Path(args.theme_json).read_text(encoding="utf-8"))
    missing = [s for s in SLOTS if s not in plan.get("colors", {})]
    if missing:
        sys.exit(f"colors に不足: {', '.join(missing)}")
    for w in check_palette({k: v.lstrip('#') for k, v in plan["colors"].items()}):
        print(f"[配色] {w}")
    prs = build(plan)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"written: {out}  layouts: {', '.join(l.name for l in prs.slide_layouts)}")


if __name__ == "__main__":
    main()
