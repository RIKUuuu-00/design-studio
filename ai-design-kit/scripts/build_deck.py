#!/usr/bin/env python3
"""骨子（deck.json）をテンプレのレイアウトに流し込んで .pptx を作る。

    python build_deck.py decks/q3/deck.json --out decks/q3/deck.pptx [--template path] [--report report.json]

方針: テンプレを開き、既存スライドを外し、レイアウトからスライドを起こしてプレースホルダーに流し込む。
マスター・テーマ・レイアウトはテンプレのまま残るので、PowerPoint でそのまま編集できる。
図表は PowerPoint ネイティブの表・グラフで作る。色はテーマ色（accent1 など）で指定する。
テンプレに無い構図は「パターン」（kpi / cards / steps / timeline / compare / matrix / chart_takeaway /
statement / agenda）で組む。文字は実フォントで測って枠に収める（タイトルは縮めず、報告だけする）。

仕様は deck-lab Skill の deck-spec.md を参照。
"""

from __future__ import annotations

import argparse
import copy
import json
import sys
from pathlib import Path

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.shapes.placeholder import ChartPlaceholder, TablePlaceholder

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib.draw import (  # noqa: E402
    NARROW_TEXT_SHAPES,
    SHAPES,
    SpecError,
    TypeScale,
    add_chart,
    add_image,
    add_table,
    fill_shape,
    line_shape,
    write_text,
)
from lib.patterns import PATTERNS, draw_pattern  # noqa: E402
from lib.pptx_common import layout_placeholder_keys, open_presentation, remove_all_slides  # noqa: E402
from lib.textfit import apply_scale, fit_scale, measure_shape  # noqa: E402

BODY_MIN_SCALE = 0.85


# ---------------------------------------------------------------- 位置

def content_area(prs, keyed: list[dict]) -> tuple[int, int, int, int]:
    """タイトルの下からフッターの上までの、図表・パターンを置く領域。"""
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    title = next((it for it in keyed if it["role"] == "title"), None)
    feet = [it["box"] for it in keyed if it["footer"] and it["box"][3] > 0]
    gap = int(sh * 0.035)
    if title:
        tx, ty, tw, th = title["box"]
        left, right, top = tx, tx + tw, ty + th + int(sh * 0.06)
    else:
        left, right, top = int(sw * 0.06), int(sw * 0.94), int(sh * 0.1)
    bottom = min([f[1] for f in feet if f[1] > top] + [int(sh * 0.92)]) - gap
    # タイトル枠の内部余白ぶん内側に寄せ、タイトルの文字と左端をそろえる
    inset = int(0.1 * 914400)
    return left + inset, top, (right - left) - 2 * inset, bottom - top


def resolve_box(prs, by_key: dict, keyed: list[dict], item: dict) -> tuple[int, int, int, int]:
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    if "box" in item:
        x, y, w, h = item["box"]
        return int(sw * x / 100), int(sh * y / 100), int(sw * w / 100), int(sh * h / 100)
    key = item.get("area", "content")
    if key == "content":
        return content_area(prs, keyed)
    if key not in by_key:
        raise SpecError(f"area '{key}' はこのレイアウトに無い。使えるキー: content, {', '.join(by_key)}")
    return by_key[key]["box"]


# ---------------------------------------------------------------- add

def add_free_shape(slide, prs, by_key, keyed, item: dict, base_dir: Path, ts: TypeScale, notes: list[str]):
    shapes = slide.shapes
    kind = item.get("type", "shape")
    if kind in PATTERNS:
        box = resolve_box(prs, by_key, keyed, item)
        notes.extend(draw_pattern(prs, slide, box, item))
        return None
    if kind == "grid":
        return add_grid(slide, prs, by_key, keyed, item, base_dir, ts, notes)
    if kind == "line":
        sw, sh = int(prs.slide_width), int(prs.slide_height)
        (x1, y1), (x2, y2) = item["from"], item["to"]
        from lib.draw import line  # noqa: PLC0415

        return line(shapes, sw * x1 / 100, sh * y1 / 100, sw * x2 / 100, sh * y2 / 100, item.get("color", "text2"), item.get("width", 1.5))
    box = resolve_box(prs, by_key, keyed, item)
    if kind == "table":
        return add_table(shapes, box, item, ts=ts)
    if kind == "chart":
        return add_chart(shapes, box, item, ts=ts)
    if kind == "image":
        return add_image(shapes, box, (base_dir / item["path"]).resolve())
    if kind == "text":
        tb = shapes.add_textbox(*box)
        tb.text_frame.word_wrap = True
        write_text(tb.text_frame, item.get("text", ""), item)
        return tb
    if kind == "shape":
        name = item.get("shape", "rect")
        if name not in SHAPES:
            raise SpecError(f"図形 '{name}' は未対応。使える図形: {', '.join(SHAPES)}")
        shp = shapes.add_shape(SHAPES[name], *box)
        shp.shadow.inherit = False
        fill_shape(shp, item.get("fill", "accent1"))
        line_shape(shp, item.get("line", "none"))
        if item.get("text") is not None:
            # 矢羽根・矢印は内側の文字領域が狭く、折り返すと 1 文字ずつ縦に並ぶ。改行は \n で明示させる
            shp.text_frame.word_wrap = name not in NARROW_TEXT_SHAPES
            opts = {"color": item.get("text_color", "bg1"), "align": item.get("align", "center"), "valign": item.get("valign", "middle")}
            for k in ("size", "bold"):
                if k in item:
                    opts[k] = item[k]
            write_text(shp.text_frame, item["text"], opts)
        return shp
    raise SpecError(
        f"add の type '{kind}' は未対応（パターン: {', '.join(PATTERNS)} / 部品: text, shape, line, table, chart, image, grid）"
    )


def add_grid(slide, prs, by_key, keyed, item: dict, base_dir: Path, ts, notes):
    """box / area を cols × rows に等分し、items を順に置く。"""
    x, y, w, h = resolve_box(prs, by_key, keyed, item)
    cols, rows = int(item.get("cols", len(item["items"]))), int(item.get("rows", 1))
    gap = int(prs.slide_width * item.get("gap", 2) / 100)
    cw = (w - gap * (cols - 1)) // cols
    ch = (h - gap * (rows - 1)) // rows
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    for i, child in enumerate(item["items"]):
        r, c = divmod(i, cols)
        if r >= rows:
            raise SpecError(f"grid の items が {cols}×{rows} を超えている")
        cx, cy = x + c * (cw + gap), y + r * (ch + gap)
        child = {**child, "box": [cx / sw * 100, cy / sh * 100, cw / sw * 100, ch / sh * 100]}
        child.pop("area", None)
        add_free_shape(slide, prs, by_key, keyed, child, base_dir, ts, notes)


# ---------------------------------------------------------------- レイアウト・プレースホルダー

def all_layouts(prs) -> list[tuple[str, object]]:
    out = []
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            out.append((f"L{len(out) + 1:02d}", layout))
    return out


def find_layout(layouts, name: str):
    for lid, layout in layouts:
        if name == lid or name == layout.name:
            return layout
    for lid, layout in layouts:
        if name.strip().lower() == layout.name.strip().lower():
            return layout
    raise SpecError(f"レイアウト '{name}' が無い。使えるレイアウト: " + ", ".join(f"{lid} {l.name}" for lid, l in layouts))


def as_kind(ph, cls):
    if isinstance(ph, cls):
        return ph
    if ph.placeholder_format.type in (PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.CHART, PP_PLACEHOLDER.TABLE):
        return cls(ph._element, ph._parent)
    # 画像枠などに表・グラフを指定された場合は枠を外し、同じ位置に普通の図形として置く
    ph._element.getparent().remove(ph._element)
    return None


def copy_footer_placeholder(slide, layout, ph_type, text: str | None = None) -> bool:
    for lph in layout.placeholders:
        if lph.placeholder_format.type == ph_type:
            el = copy.deepcopy(lph._element)
            # 図形 ID はスライド内で一意でないと PowerPoint が修復を求めるため振り直す
            ids = [int(v) for v in slide.shapes._spTree.xpath("//p:cNvPr/@id")]
            el.xpath("./*[1]/p:cNvPr")[0].set("id", str(max(ids + [1]) + 1))
            slide.shapes._spTree.append(el)
            if text is not None:
                new = slide.placeholders[lph.placeholder_format.idx]
                write_text(new.text_frame, text)
            return True
    return False


def fit_placeholder(ph, slide, role: str, notes: list[str], policy: str) -> None:
    """本文は最大 15% まで縮めて収める。タイトルは縮めず、溢れ・1〜2 文字だけの最終行を報告する。"""
    r = measure_shape(ph, slide)
    if r is None:
        return
    label = "タイトル" if role == "title" else ph.name
    if r["widows"]:
        notes.append(f"{label}: 最終行が「{r['widows'][0]}」だけになる。言い回しを変えて行末をそろえる")
    if r["ratio"] <= 1.0:
        return
    if role == "title" or policy == "none":
        notes.append(f"{label}: 枠の {r['ratio']:.0%}（{r['line_count']} 行）。短くする")
        return
    s = fit_scale(ph, slide, BODY_MIN_SCALE)
    if s is None:
        apply_scale(ph, slide, BODY_MIN_SCALE)
        notes.append(f"{label}: {int(BODY_MIN_SCALE * 100)}% に縮めても収まらない。文字を減らすかスライドを分ける")
    else:
        apply_scale(ph, slide, s)
        notes.append(f"{label}: 文字を {int(s * 100)}% に縮めて収めた")


def build(spec: dict, template: Path, base_dir: Path):
    prs = open_presentation(template)
    if not spec.get("keep_template_slides", False):
        remove_all_slides(prs)
    layouts = all_layouts(prs)
    ts = TypeScale(int(prs.slide_height), spec.get("base_size"), int(prs.slide_width))
    fit_policy = spec.get("fit", "shrink")
    report = []
    for s_no, s in enumerate(spec["slides"], start=1):
        try:
            layout = find_layout(layouts, s["layout"])
            slide = prs.slides.add_slide(layout)
            keyed = layout_placeholder_keys(layout)
            by_key = {it["key"]: it for it in keyed}
            filled: dict[int, str] = {}
            notes: list[str] = []
            for key, value in (s.get("fill") or {}).items():
                if key.startswith("#"):
                    idx = int(key[1:])
                    role = next((it["role"] for it in keyed if it["idx"] == idx), "body")
                else:
                    match = by_key.get(key) or next((it for it in keyed if it["name"] == key), None)
                    if match is None:
                        raise SpecError(
                            f"キー '{key}' はレイアウト '{layout.name}' に無い。使えるキー: {', '.join(k for k, it in by_key.items() if not it['footer'])}"
                        )
                    idx, role = match["idx"], match["role"]
                try:
                    ph = slide.placeholders[idx]
                except KeyError as exc:
                    raise SpecError(f"idx {idx} のプレースホルダーがスライドに無い") from exc
                box = (int(ph.left), int(ph.top), int(ph.width), int(ph.height))
                filled[idx] = role
                if isinstance(value, dict) and "table" in value:
                    # 本文（OBJECT）枠にも表を入れられるよう TablePlaceholder として扱う
                    add_table(slide.shapes, box, value["table"], as_kind(ph, TablePlaceholder), ts=ts)
                elif isinstance(value, dict) and "chart" in value:
                    add_chart(slide.shapes, box, value["chart"], as_kind(ph, ChartPlaceholder), ts=ts)
                elif isinstance(value, dict) and "image" in value:
                    if hasattr(ph, "insert_picture"):
                        add_image(slide.shapes, box, (base_dir / value["image"]).resolve(), ph)
                    else:
                        ph._element.getparent().remove(ph._element)
                        add_image(slide.shapes, box, (base_dir / value["image"]).resolve())
                elif isinstance(value, dict) and value.get("pattern"):
                    # 本文枠の位置にパターンを組む（枠は外す）
                    ph._element.getparent().remove(ph._element)
                    notes.extend(draw_pattern(prs, slide, box, {**value, "type": value["pattern"]}))
                else:
                    write_text(ph.text_frame, value)
            # insert_* 系は元の要素を置き換えるので、ここで改めて空き枠を集める
            removed = []
            if not s.get("keep_empty", False):
                for ph in list(slide.placeholders):
                    idx = ph.placeholder_format.idx
                    if idx in filled:
                        continue
                    if ph.has_text_frame and ph.text_frame.text.strip():
                        continue
                    removed.append(ph.name)
                    ph._element.getparent().remove(ph._element)
            # 文字を測って収める
            for ph in list(slide.placeholders):
                role = filled.get(ph.placeholder_format.idx)
                if role and getattr(ph, "has_text_frame", False):
                    fit_placeholder(ph, slide, role, notes, s.get("fit", fit_policy))
            for item in s.get("add") or []:
                add_free_shape(slide, prs, by_key, keyed, item, base_dir, ts, notes)
            page_numbers = s.get("page_number", spec.get("page_numbers", True))
            if page_numbers:
                copy_footer_placeholder(slide, layout, PP_PLACEHOLDER.SLIDE_NUMBER)
            footer = s.get("footer", spec.get("footer"))
            if footer:
                copy_footer_placeholder(slide, layout, PP_PLACEHOLDER.FOOTER, footer)
            if s.get("notes"):
                slide.notes_slide.notes_text_frame.text = s["notes"]
            report.append({"no": s_no, "layout": layout.name, "removed_empty": removed, "notes": notes})
        except SpecError as exc:
            raise SpecError(f"スライド {s_no}: {exc}") from exc
    if spec.get("title"):
        prs.core_properties.title = spec["title"]
    return prs, report


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", help="deck.json")
    ap.add_argument("--out", required=True)
    ap.add_argument("--template", help="spec の template より優先")
    ap.add_argument("--report", help="スライドごとの調整・注意を JSON で書き出す")
    args = ap.parse_args()

    spec_path = Path(args.spec)
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    template = Path(args.template or spec.get("template") or "")
    if not template.is_file():
        # spec からの相対パスも試す
        alt = (spec_path.parent / template).resolve()
        if alt.is_file():
            template = alt
        else:
            sys.exit(f"テンプレが見つからない: {template}")
    try:
        prs, report = build(spec, template, spec_path.parent)
    except SpecError as exc:
        sys.exit(f"[spec エラー] {exc}")
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    if args.report:
        Path(args.report).write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"written: {out}  slides: {len(report)}")
    for r in report:
        print(f"  {r['no']:>2}. {r['layout']}")
        for n in r["notes"]:
            print(f"      - {n}")


if __name__ == "__main__":
    main()
