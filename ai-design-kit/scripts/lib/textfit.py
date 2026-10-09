"""実フォントの字幅で、テキスト枠に収まるかを計算する。

- フォントは fc-match（Linux/macOS）か Windows の既定フォルダから探す。見つからなければ全角=1em の概算に落とす
- 改行は日本語の禁則（行頭に句読点・閉じ括弧を置かない）と英単語の分割禁止を考慮する
- 文字サイズ・行間・段落前後の間隔・インデント・内部余白は、スライド → レイアウト → マスターの継承を辿って解決する
"""

from __future__ import annotations

import math
import os
import shutil
import subprocess
import unicodedata
from functools import lru_cache

from .pptx_common import DEFAULT_INSET_LR, DEFAULT_INSET_TB, EMU_PER_PT, NS, placeholder_role

# 行頭に置かない文字（句読点・閉じ括弧・長音・小書き仮名など）
NO_LINE_START = set("、。，．,.)）」』】〕〉》｝]!！?？:：;；・ー～ぁぃぅぇぉっゃゅょゎァィゥェォッャュョヮヵヶ％%")
# 行末に置かない文字（開き括弧）
NO_LINE_END = set("(（「『【〔〈《｛[")

WINDOWS_FONTS = {
    "meiryo": "meiryo.ttc", "メイリオ": "meiryo.ttc", "meiryo ui": "meiryo.ttc",
    "yu gothic": "YuGothM.ttc", "游ゴシック": "YuGothM.ttc", "yu gothic ui": "YuGothM.ttc",
    "yu mincho": "yumin.ttf", "游明朝": "yumin.ttf",
    "biz udpgothic": "BIZ-UDGothicR.ttc", "biz udpゴシック": "BIZ-UDGothicR.ttc", "biz udgothic": "BIZ-UDGothicR.ttc",
    "biz udpmincho": "BIZ-UDMinchoM.ttc", "biz udp明朝 medium": "BIZ-UDMinchoM.ttc",
    "ms gothic": "msgothic.ttc", "ｍｓ ゴシック": "msgothic.ttc", "ms pgothic": "msgothic.ttc",
    "arial": "arial.ttf", "calibri": "calibri.ttf", "segoe ui": "segoeui.ttf",
}


# ---------------------------------------------------------------- フォント

@lru_cache(maxsize=64)
def font_file(family: str, bold: bool = False) -> str | None:
    if not family:
        return None
    if shutil.which("fc-match"):
        pattern = f"{family}:lang=ja" + (":bold" if bold else "")
        try:
            out = subprocess.run(["fc-match", "-f", "%{file}|%{family}", pattern], capture_output=True, text=True, timeout=10).stdout
            path, _, fam = out.partition("|")
            # fc-match は無いフォントにも代替を返す。代替でも実際に描画に使われるのはそれなので採用する
            if path and os.path.exists(path):
                return path
        except Exception:  # noqa: BLE001
            pass
    windir = os.environ.get("WINDIR")
    if windir:
        name = WINDOWS_FONTS.get(family.strip().lower())
        for base in (os.path.join(windir, "Fonts"), os.path.expandvars(r"%LOCALAPPDATA%\Microsoft\Windows\Fonts")):
            if name and os.path.exists(os.path.join(base, name)):
                return os.path.join(base, name)
    return None


@lru_cache(maxsize=64)
def _pil_font(path: str):
    from PIL import ImageFont

    return ImageFont.truetype(path, 100)


class Measurer:
    """1 フォント分の字幅・行の高さを測る。PIL かフォントが無ければ概算する。"""

    def __init__(self, family: str | None, bold: bool = False):
        self.family = family
        self.path = None
        self.font = None
        try:
            self.path = font_file(family or "", bold) or font_file("sans-serif", bold)
            if self.path:
                self.font = _pil_font(self.path)
        except Exception:  # noqa: BLE001
            self.font = None

    @property
    def exact(self) -> bool:
        return self.font is not None

    def width_em(self, text: str) -> float:
        if not text:
            return 0.0
        if self.font is not None:
            return self.font.getlength(text) / 100
        return sum(1.0 if unicodedata.east_asian_width(c) in ("F", "W", "A") else 0.55 for c in text)

    def line_em(self) -> float:
        """行送り（1.0 行間）を em で返す。"""
        if self.font is not None:
            ascent, descent = self.font.getmetrics()
            return max((ascent + descent) / 100, 1.0)
        return 1.2


# ---------------------------------------------------------------- 改行

def _units(text: str) -> list[str]:
    """改行できる単位に分ける（英数字の連続は 1 単位、全角は 1 文字ずつ）。"""
    units: list[str] = []
    buf = ""
    for ch in text:
        wide = unicodedata.east_asian_width(ch) in ("F", "W", "A")
        if wide or ch == " ":
            if buf:
                units.append(buf)
                buf = ""
            units.append(ch)
        else:
            buf += ch
    if buf:
        units.append(buf)
    return units


def wrap(text: str, max_em: float, m: Measurer) -> list[str]:
    if max_em <= 0:
        return [text]
    lines: list[str] = []
    cur = ""
    for u in _units(text):
        if cur and m.width_em(cur + u) > max_em:
            if u[0] in NO_LINE_START:
                # 句読点はぶら下げる（行末に残す）
                cur += u
                continue
            if cur[-1] in NO_LINE_END:
                lines.append(cur[:-1].rstrip())
                cur = cur[-1] + u
                continue
            lines.append(cur.rstrip())
            cur = u.lstrip() if u == " " else u
            # 1 単位で枠より長い（長い英単語・URL）なら文字単位で割る
            while m.width_em(cur) > max_em and len(cur) > 1:
                cut = len(cur)
                while cut > 1 and m.width_em(cur[:cut]) > max_em:
                    cut -= 1
                lines.append(cur[:cut])
                cur = cur[cut:]
        else:
            cur += u
    lines.append(cur.rstrip())
    return lines


# ---------------------------------------------------------------- 継承の解決

def _chain(shape, slide):
    """段落・本文プロパティを探す順に (要素, 種類) を返す。"""
    chain = [shape._element]
    role = None
    if getattr(shape, "is_placeholder", False) and slide is not None:
        try:
            idx = shape.placeholder_format.idx
            role = placeholder_role(shape)
            layout = slide.slide_layout
            lph = next((p for p in layout.placeholders if p.placeholder_format.idx == idx), None)
            if lph is None:
                lph = next((p for p in layout.placeholders if placeholder_role(p) == role), None)
            if lph is not None:
                chain.append(lph._element)
            master = layout.slide_master
            mph = next((p for p in master.placeholders if placeholder_role(p) == role), None)
            if mph is None and role in ("subtitle",):
                mph = next((p for p in master.placeholders if placeholder_role(p) == "body"), None)
            if mph is not None:
                chain.append(mph._element)
            tx = master._element.find("p:txStyles", NS)
            if tx is not None:
                style = {"title": "titleStyle"}.get(role, "bodyStyle")
                chain.append(tx.find(f"p:{style}", NS))
        except Exception:  # noqa: BLE001
            pass
    elif slide is not None:
        try:
            tx = slide.slide_layout.slide_master._element.find("p:txStyles", NS)
            if tx is not None:
                chain.append(tx.find("p:otherStyle", NS))
        except Exception:  # noqa: BLE001
            pass
    return [c for c in chain if c is not None], role


def _lvl(el, level: int):
    """要素（sp か txStyles の子）から lvlNpPr を探す。"""
    if el.tag.endswith("Style") and not el.tag.endswith("lstStyle"):
        return el.find(f"a:lvl{level}pPr", NS)
    return el.find(f".//a:lstStyle/a:lvl{level}pPr", NS)


def para_props(shape, slide, para, level: int) -> dict:
    """段落の文字サイズ・太字・行間・段落前後・左余白を解決する。"""
    chain, role = _chain(shape, slide)
    props = {"size": None, "bold": None, "line": None, "before": None, "after": None, "marL": None, "indent": None, "font_ea": None, "font_latin": None}
    p_pr = para._p.find("a:pPr", NS)
    sources = []
    if p_pr is not None:
        sources.append(p_pr)
    for el in chain:
        lv = _lvl(el, level)
        if lv is not None:
            sources.append(lv)
    for src in sources:
        rpr = src.find("a:defRPr", NS)
        if rpr is not None:
            if props["size"] is None and rpr.get("sz"):
                props["size"] = int(rpr.get("sz")) / 100
            if props["bold"] is None and rpr.get("b") is not None:
                props["bold"] = rpr.get("b") in ("1", "true")
            for tag, key in (("ea", "font_ea"), ("latin", "font_latin")):
                f = rpr.find(f"a:{tag}", NS)
                if props[key] is None and f is not None and f.get("typeface"):
                    props[key] = f.get("typeface")
        if props["line"] is None:
            ln = src.find("a:lnSpc/a:spcPct", NS)
            if ln is not None:
                props["line"] = int(ln.get("val")) / 100000
            ln_pts = src.find("a:lnSpc/a:spcPts", NS)
            if ln_pts is not None:
                props["line"] = ("pts", int(ln_pts.get("val")) / 100)
        for tag, key in (("spcBef", "before"), ("spcAft", "after")):
            if props[key] is None:
                pct = src.find(f"a:{tag}/a:spcPct", NS)
                pts = src.find(f"a:{tag}/a:spcPts", NS)
                if pct is not None:
                    props[key] = ("pct", int(pct.get("val")) / 100000)
                elif pts is not None:
                    props[key] = ("pts", int(pts.get("val")) / 100)
        for key in ("marL", "indent"):
            if props[key] is None and src.get(key) is not None:
                props[key] = int(src.get(key))
    # 段落内の run の明示サイズ・太字が優先
    for r in para.runs:
        if r.font.size is not None:
            props["size"] = r.font.size.pt
        if r.font.bold is not None:
            props["bold"] = r.font.bold
        break
    if props["size"] is None:
        props["size"] = 18.0
    props["role"] = role
    return props


def body_insets(shape, slide) -> tuple[int, int, int, int]:
    chain, _ = _chain(shape, slide)
    vals = {}
    for el in chain:
        bp = el.find(".//a:bodyPr", NS)
        if bp is None:
            continue
        for k in ("lIns", "tIns", "rIns", "bIns"):
            if k not in vals and bp.get(k) is not None:
                vals[k] = int(bp.get(k))
    return (
        vals.get("lIns", DEFAULT_INSET_LR),
        vals.get("tIns", DEFAULT_INSET_TB),
        vals.get("rIns", DEFAULT_INSET_LR),
        vals.get("bIns", DEFAULT_INSET_TB),
    )


def autofit_scale(shape) -> float:
    na = shape._element.find(".//a:bodyPr/a:normAutofit", NS)
    if na is not None and na.get("fontScale"):
        return int(na.get("fontScale")) / 100000
    return 1.0


# ---------------------------------------------------------------- 測定

def theme_fonts(slide) -> tuple[str | None, str | None]:
    """(見出しの ea/latin フォント, 本文の ea/latin フォント)。"""
    from .pptx_common import theme_info  # noqa: PLC0415

    try:
        prs_part = slide.part.package.presentation_part
        t = theme_info(prs_part.presentation)
    except Exception:  # noqa: BLE001
        return None, None
    major = t["fonts"]["major"].get("ea") or t["fonts"]["major"].get("latin")
    minor = t["fonts"]["minor"].get("ea") or t["fonts"]["minor"].get("latin")
    return major, minor


def measure_shape(shape, slide, scale: float = 1.0) -> dict | None:
    """テキスト枠の必要高さ・行情報を返す。scale は文字サイズの倍率（縮小の試算用）。"""
    if not getattr(shape, "has_text_frame", False) or shape.width is None or shape.height is None:
        return None
    tf = shape.text_frame
    if not tf.text.strip():
        return None
    l_ins, t_ins, r_ins, b_ins = body_insets(shape, slide)
    inner_w = max(int(shape.width) - l_ins - r_ins, 1)
    inner_h = max(int(shape.height) - t_ins - b_ins, 1)
    wrap_on = tf.word_wrap is not False
    major, minor = theme_fonts(slide)
    total = 0.0
    lines_info = []
    widest = 0.0
    exact = True
    max_size = 0.0
    fscale = autofit_scale(shape) * scale
    for i, para in enumerate(tf.paragraphs):
        level = min(para.level + 1, 9)
        pp = para_props(shape, slide, para, level)
        size = pp["size"] * fscale
        max_size = max(max_size, size)
        family = pp["font_ea"] or pp["font_latin"]
        if family is None or family.startswith("+mj"):
            family = major if (pp["role"] == "title" or family) else minor
        elif family.startswith("+mn"):
            family = minor
        if family is None or family.startswith("+"):
            family = major if pp["role"] == "title" else minor
        m = Measurer(family, bool(pp["bold"]))
        exact = exact and m.exact
        size_emu = size * EMU_PER_PT
        mar = (pp["marL"] or 0)
        avail_em = max((inner_w - mar) / size_emu, 1)
        text = "".join(r.text for r in para.runs).replace("\v", "\n")
        para_lines = []
        for chunk in text.split("\n") or [""]:
            para_lines.extend(wrap(chunk, avail_em, m) if wrap_on else [chunk])
        widest = max([widest] + [m.width_em(ln) * size_emu for ln in para_lines])
        line = pp["line"]
        if isinstance(line, tuple):
            line_h = line[1] * EMU_PER_PT
        else:
            line_h = size_emu * m.line_em() * (line or 1.0)
        h = line_h * len(para_lines)
        for key in ("before", "after"):
            sp = pp[key]
            if sp and not (key == "before" and i == 0):
                h += sp[1] * EMU_PER_PT if sp[0] == "pts" else sp[1] * size_emu
        total += h
        lines_info.append({"lines": para_lines, "size": round(size, 1)})
    widows = [
        p["lines"][-1] for p in lines_info
        if len(p["lines"]) >= 2 and 0 < len(p["lines"][-1].strip()) <= 2
    ]
    return {
        "ratio": round(total / inner_h, 2),
        "needed_emu": int(total),
        "inner_h": inner_h,
        "inner_w": inner_w,
        "widest_ratio": round(widest / inner_w, 2) if inner_w else 0,
        "max_size_pt": round(max_size, 1),
        "line_count": sum(len(p["lines"]) for p in lines_info),
        "paragraphs": lines_info,
        "widows": widows,
        "exact": exact,
        "wrap": wrap_on,
    }


def fit_scale(shape, slide, min_scale: float = 0.8, step: float = 0.05) -> float | None:
    """枠に収まる最大の倍率（1.0 から step 刻み）。min_scale でも収まらなければ None。"""
    s = 1.0
    while s >= min_scale - 1e-9:
        r = measure_shape(shape, slide, s)
        if r is None or (r["ratio"] <= 1.0 and (r["wrap"] or r["widest_ratio"] <= 1.0)):
            return round(s, 2)
        s -= step
    return None


def apply_scale(shape, slide, scale: float) -> None:
    """全段落の run に解決済みサイズ × scale を明示する。"""
    for para in shape.text_frame.paragraphs:
        level = min(para.level + 1, 9)
        size = para_props(shape, slide, para, level)["size"]
        from pptx.util import Pt  # noqa: PLC0415

        new = Pt(max(math.floor(size * scale * 2) / 2, 1))
        for r in para.runs:
            r.font.size = new
