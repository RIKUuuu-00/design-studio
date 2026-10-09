"""deck-lab の各スクリプトが共有するヘルパー。

python-pptx だけで完結させ、外部 Skill のコードには依存しない。
"""

from __future__ import annotations

import copy
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}

EMU_PER_PT = 12700
EMU_PER_INCH = 914400

# 既定の内部余白（OOXML の既定値）
DEFAULT_INSET_LR = int(0.1 * EMU_PER_INCH)
DEFAULT_INSET_TB = int(0.05 * EMU_PER_INCH)

# スライドへ複製されない（フッター系）プレースホルダー
FOOTER_TYPES = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}

ROLE_BY_TYPE = {
    PP_PLACEHOLDER.TITLE: "title",
    PP_PLACEHOLDER.CENTER_TITLE: "title",
    PP_PLACEHOLDER.VERTICAL_TITLE: "title",
    PP_PLACEHOLDER.SUBTITLE: "subtitle",
    PP_PLACEHOLDER.BODY: "body",
    PP_PLACEHOLDER.VERTICAL_BODY: "body",
    PP_PLACEHOLDER.OBJECT: "body",
    PP_PLACEHOLDER.VERTICAL_OBJECT: "body",
    PP_PLACEHOLDER.PICTURE: "picture",
    PP_PLACEHOLDER.BITMAP: "picture",
    PP_PLACEHOLDER.CHART: "chart",
    PP_PLACEHOLDER.TABLE: "table",
    PP_PLACEHOLDER.ORG_CHART: "diagram",
    PP_PLACEHOLDER.MEDIA_CLIP: "media",
    PP_PLACEHOLDER.DATE: "date",
    PP_PLACEHOLDER.FOOTER: "footer",
    PP_PLACEHOLDER.SLIDE_NUMBER: "slide_number",
}

# テンプレの見本文・入力促し文として残りやすい文字列
LEFTOVER_PATTERNS = [
    r"lorem ipsum",
    r"click to (add|edit)",
    r"クリックして",
    r"テキストを入力",
    r"ここに.{0,12}(入力|記入|入れ)",
    r"[○〇◯]{2,}",
    r"\bX{2,}\b",
    r"\bxxx+\b",
    r"\bTBD\b",
    r"ダミー",
    r"サンプルテキスト",
    r"(title|subtitle|body) here",
]


# ---------------------------------------------------------------- 入出力

def open_presentation(path: str | Path) -> Presentation:
    """.pptx / .potx を開く。.potx は一時的に .pptx の content type に書き換える。"""
    path = Path(path)
    if path.suffix.lower() != ".potx":
        return Presentation(str(path))
    tmp = Path(tempfile.mkdtemp(prefix="decklab-")) / (path.stem + ".pptx")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"presentationml.template.main+xml",
                    b"presentationml.presentation.main+xml",
                )
            dst.writestr(item, data)
    return Presentation(str(tmp))


def remove_all_slides(prs: Presentation) -> int:
    """既存スライドを全て外す（マスター・レイアウト・テーマは残る）。"""
    sld_id_lst = prs.slides._sldIdLst
    removed = 0
    for sld_id in list(sld_id_lst):
        prs.part.drop_rel(sld_id.rId)
        sld_id_lst.remove(sld_id)
        removed += 1
    return removed


# ---------------------------------------------------------------- 幾何

def emu_to_pct(value: int, total: int) -> float:
    return round(value / total * 100, 1) if total else 0.0


def shape_box(shape) -> tuple[int, int, int, int] | None:
    if shape.left is None or shape.top is None or shape.width is None or shape.height is None:
        return None
    return int(shape.left), int(shape.top), int(shape.width), int(shape.height)


def intersection_area(a, b) -> int:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    w = min(ax + aw, bx + bw) - max(ax, bx)
    h = min(ay + ah, by + bh) - max(ay, by)
    return w * h if w > 0 and h > 0 else 0


# ---------------------------------------------------------------- プレースホルダー

def placeholder_role(ph) -> str:
    try:
        return ROLE_BY_TYPE.get(ph.placeholder_format.type, "body")
    except Exception:  # noqa: BLE001 - 型不明のプレースホルダーは本文扱い
        return "body"


def layout_placeholder_keys(layout) -> list[dict]:
    """レイアウトのプレースホルダーに安定したキー（title, body, body2 ...）を振る。

    同じ役割が複数ある場合は「上→下、左→右」の順に番号を付ける。
    """
    items = []
    for ph in layout.placeholders:
        box = shape_box(ph) or (0, 0, 0, 0)
        items.append({"ph": ph, "role": placeholder_role(ph), "box": box})
    items.sort(key=lambda it: (round(it["box"][1] / EMU_PER_INCH, 1), it["box"][0]))
    counts: dict[str, int] = {}
    for it in items:
        counts[it["role"]] = counts.get(it["role"], 0) + 1
        n = counts[it["role"]]
        it["key"] = it["role"] if n == 1 else f"{it['role']}{n}"
        it["idx"] = it["ph"].placeholder_format.idx
        it["name"] = it["ph"].name
        it["type"] = str(it["ph"].placeholder_format.type).split(".")[-1].split(" ")[0]
        it["footer"] = it["ph"].placeholder_format.type in FOOTER_TYPES
    return items


# ---------------------------------------------------------------- テーマ

def theme_info(prs: Presentation) -> dict:
    """最初のスライドマスターのテーマから色とフォントを取り出す。"""
    master = prs.slide_masters[0]
    theme_part = None
    for rel in master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            theme_part = rel.target_part
            break
    info = {"name": "", "colors": {}, "fonts": {"major": {}, "minor": {}}}
    if theme_part is None:
        return info
    from lxml import etree

    root = etree.fromstring(theme_part.blob)
    info["name"] = root.get("name", "")
    scheme = root.find(".//a:clrScheme", NS)
    if scheme is not None:
        for child in scheme:
            tag = etree.QName(child).localname
            val = None
            for c in child:
                val = c.get("lastClr") or c.get("val")
            info["colors"][tag] = (val or "").upper()
    for kind in ("major", "minor"):
        font = root.find(f".//a:{kind}Font", NS)
        if font is None:
            continue
        for script in ("latin", "ea", "cs"):
            el = font.find(f"a:{script}", NS)
            if el is not None and el.get("typeface"):
                info["fonts"][kind][script] = el.get("typeface")
        jpan = font.find("a:font[@script='Jpan']", NS)
        if jpan is not None and not info["fonts"][kind].get("ea"):
            info["fonts"][kind]["ea"] = jpan.get("typeface")
    return info


def installed_font_families() -> set[str] | None:
    """fc-list で導入済みフォント名を返す。fc-list が無い環境では None。"""
    exe = shutil.which("fc-list")
    if not exe:
        return None
    try:
        out = subprocess.run([exe, ":", "family"], capture_output=True, text=True, timeout=30).stdout
    except Exception:  # noqa: BLE001
        return None
    names: set[str] = set()
    for line in out.splitlines():
        for name in line.split(","):
            names.add(name.strip().lower())
    return names


def missing_theme_fonts(theme: dict) -> list[str]:
    installed = installed_font_families()
    if installed is None:
        return []
    wanted = set()
    for kind in ("major", "minor"):
        for script in ("latin", "ea"):
            face = theme["fonts"].get(kind, {}).get(script)
            if face and not face.startswith("+"):
                wanted.add(face)
    return sorted(f for f in wanted if f.lower() not in installed)


# ---------------------------------------------------------------- 文字サイズ推定

def _sz_from(el, level: int = 1) -> float | None:
    """txBody / lstStyle を含む要素から lvl の既定サイズ（pt）を探す。"""
    if el is None:
        return None
    for path in (
        f".//a:lstStyle/a:lvl{level}pPr/a:defRPr",
        ".//a:p/a:pPr/a:defRPr",
        ".//a:p/a:r/a:rPr",
        ".//a:p/a:endParaRPr",
    ):
        node = el.find(path, NS)
        if node is not None and node.get("sz"):
            return int(node.get("sz")) / 100
    return None


def layout_placeholder_font_size(layout, ph, level: int = 1) -> float | None:
    """レイアウト上のプレースホルダー → マスター → txStyles の順で既定サイズ（pt）を探す。"""
    size = _sz_from(ph._element, level)
    if size:
        return size
    master = layout.slide_master
    role = placeholder_role(ph)
    for mph in master.placeholders:
        if placeholder_role(mph) == role:
            size = _sz_from(mph._element, level)
            if size:
                return size
    tx = master._element.find("p:txStyles", NS)
    if tx is not None:
        style = "titleStyle" if role == "title" else "bodyStyle"
        node = tx.find(f"p:{style}/a:lvl{level}pPr/a:defRPr", NS)
        if node is not None and node.get("sz"):
            return int(node.get("sz")) / 100
    return None


def chars_capacity(width: int, height: int, size_pt: float, line_spacing: float = 1.2) -> tuple[int, int]:
    """枠に入る（1行あたり全角文字数, 行数）の目安。"""
    size_emu = size_pt * EMU_PER_PT
    inner_w = max(width - 2 * DEFAULT_INSET_LR, 1)
    inner_h = max(height - 2 * DEFAULT_INSET_TB, 1)
    return max(int(inner_w / size_emu), 1), max(int(inner_h / (size_emu * line_spacing)), 1)


# ---------------------------------------------------------------- テキスト判定

def leftover_hits(text: str, extra: list[str] | None = None) -> list[str]:
    hits = []
    for pat in LEFTOVER_PATTERNS + (extra or []):
        if re.search(pat, text, flags=re.IGNORECASE):
            hits.append(pat)
    return hits


def clone_element(el):
    return copy.deepcopy(el)


def find_soffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        exe = shutil.which(name)
        if exe:
            return exe
    for cand in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if os.path.exists(cand):
            return cand
    return None

