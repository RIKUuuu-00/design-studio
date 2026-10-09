#!/usr/bin/env python3
"""社内テンプレを解析し、Claude が使うレイアウトカタログを作る。

    python register_template.py design-system/slides/template.pptx [--out design-system/slides] [--no-render]

出力（--out 配下）:
- layouts.json   … build_deck.py が使う機械可読カタログ（レイアウト名・プレースホルダーキー・位置）
- layouts.md     … Claude が骨子を組むときに読むカタログ。各レイアウトの「用途」は Claude が埋める
- thumbs/        … レイアウトごとのサムネ（プレースホルダーにキー名を入れた見本）と一覧 sheet.png
- 標準出力       … テンプレの構造チェック結果（直すべき箇所）

再登録しても layouts.md に書いた「用途」はレイアウト名で引き継ぐ。
"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter

from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lib.pptx_common import (  # noqa: E402
    EMU_PER_INCH,
    chars_capacity,
    emu_to_pct,
    intersection_area,
    layout_placeholder_font_size,
    layout_placeholder_keys,
    leftover_hits,
    missing_theme_fonts,
    open_presentation,
    remove_all_slides,
    theme_info,
)

GENERIC_LAYOUT_NAMES = re.compile(
    r"^(\d+_)?(custom layout|ユーザー設定レイアウト|レイアウト\s*\d*|layout\s*\d*|slide layout\s*\d*)$",
    re.IGNORECASE,
)
PURPOSE_RE = re.compile(r"^### L(\d+) (.+?)\s*$")
PURPOSE_LINE_RE = re.compile(r"^- \*\*用途\*\*: (.*)$")


def read_purposes(md_path: Path) -> dict[str, str]:
    if not md_path.exists():
        return {}
    purposes: dict[str, str] = {}
    current = None
    for line in md_path.read_text(encoding="utf-8").splitlines():
        m = PURPOSE_RE.match(line)
        if m:
            current = m.group(2)
            continue
        m = PURPOSE_LINE_RE.match(line)
        if m and current:
            if m.group(1).strip() and not m.group(1).startswith("TODO"):
                purposes[current] = m.group(1).strip()
            current = None
    return purposes


def analyze(prs, template_path: Path) -> tuple[dict, list[dict]]:
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    theme = theme_info(prs)
    issues: list[dict] = []
    layouts = []
    name_counts = Counter()
    for m_i, master in enumerate(prs.slide_masters):
        for l_i, layout in enumerate(master.slide_layouts):
            name_counts[layout.name] += 1
            phs = layout_placeholder_keys(layout)
            entry = {
                "id": f"L{len(layouts) + 1:02d}",
                "name": layout.name,
                "master": m_i,
                "index": l_i,
                "placeholders": [],
                "static_texts": [],
            }
            for it in phs:
                x, y, w, h = it["box"]
                size = layout_placeholder_font_size(layout, it["ph"]) or 18.0
                cpl, lines = chars_capacity(w, h, size)
                entry["placeholders"].append(
                    {
                        "key": it["key"],
                        "idx": it["idx"],
                        "type": it["type"],
                        "name": it["name"],
                        "footer": it["footer"],
                        "box_pct": [emu_to_pct(x, sw), emu_to_pct(y, sh), emu_to_pct(w, sw), emu_to_pct(h, sh)],
                        "font_pt": size,
                        "capacity": {"chars_per_line": cpl, "lines": lines},
                    }
                )
                if x < 0 or y < 0 or x + w > sw or y + h > sh:
                    issues.append(_issue("warn", f"{layout.name}", f"プレースホルダー「{it['name']}」がスライド外にはみ出している"))
            # レイアウト上の非プレースホルダーのテキスト（全スライドに出る）
            for shp in layout.shapes:
                if shp.is_placeholder or not getattr(shp, "has_text_frame", False):
                    continue
                text = shp.text_frame.text.strip()
                if text:
                    entry["static_texts"].append(text[:60])
                    if leftover_hits(text):
                        issues.append(_issue("warn", layout.name, f"レイアウトに見本文らしき固定テキスト「{text[:30]}」がある（全スライドに表示される）"))
            # チェック
            content = [p for p in entry["placeholders"] if not p["footer"]]
            if GENERIC_LAYOUT_NAMES.match(layout.name.strip()):
                issues.append(_issue("warn", layout.name, "レイアウト名が既定名のまま。用途が分かる名前に変えると選択精度が上がる"))
            if content and not any(p["key"] == "title" for p in content):
                issues.append(_issue("info", layout.name, "タイトルのプレースホルダーが無い（主張タイトルを入れられない）"))
            if not content:
                issues.append(_issue("info", layout.name, "流し込めるプレースホルダーが無い（白紙・装飾用として扱う）"))
            boxes = [(p, ph["box"]) for p, ph in zip(entry["placeholders"], phs) if not p["footer"]]
            for i in range(len(boxes)):
                for j in range(i + 1, len(boxes)):
                    a, b = boxes[i][1], boxes[j][1]
                    overlap = intersection_area(a, b)
                    smaller = min(a[2] * a[3], b[2] * b[3]) or 1
                    if overlap / smaller > 0.2:
                        issues.append(
                            _issue("warn", layout.name, f"プレースホルダー「{boxes[i][0]['key']}」と「{boxes[j][0]['key']}」が重なっている")
                        )
            layouts.append(entry)
    for name, n in name_counts.items():
        if n > 1:
            issues.append(_issue("warn", name, f"同名のレイアウトが {n} 個ある。build_deck は最初の 1 つしか選べない"))

    samples = []
    for s_i, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text_frame.text.strip() if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        free_text = [
            shp for shp in slide.shapes
            if not shp.is_placeholder and getattr(shp, "has_text_frame", False) and shp.text_frame.text.strip()
        ]
        samples.append({"no": s_i, "layout": slide.slide_layout.name, "title": title[:60], "free_text_boxes": len(free_text)})
        if len(free_text) >= 3:
            issues.append(
                _issue("info", f"サンプルスライド {s_i}", f"本文が {len(free_text)} 個のテキストボックスで組まれている。図解の見本として参照はできるが、流し込みはできない")
            )

    missing = missing_theme_fonts(theme)
    if missing:
        issues.append(
            _issue("warn", "テーマ", f"テーマのフォント {', '.join(missing)} がこの環境に無い。画像化の結果が実物とずれる（はみ出し判定が狂う）")
        )
    has_kind = {p["type"] for l in layouts for p in l["placeholders"]}
    if not ({"CHART", "TABLE", "OBJECT"} & has_kind):
        issues.append(_issue("info", "テンプレ全体", "グラフ・表を置けるプレースホルダーが無い。図表は本文枠の位置に直接置く"))

    catalog = {
        "template": str(template_path),
        "slide_size": {
            "width_in": round(sw / EMU_PER_INCH, 2),
            "height_in": round(sh / EMU_PER_INCH, 2),
            "aspect": "16:9" if abs(sw / sh - 16 / 9) < 0.02 else ("4:3" if abs(sw / sh - 4 / 3) < 0.02 else f"{sw / sh:.2f}"),
        },
        "theme": theme,
        "layouts": layouts,
        "sample_slides": samples,
    }
    return catalog, issues


def _issue(level: str, where: str, message: str) -> dict:
    return {"level": level, "where": where, "message": message}


def build_probe(template: Path, catalog: dict, target: Path) -> None:
    """各レイアウトを 1 枚ずつ並べ、プレースホルダーにキー名を書いた見本 deck を作る。"""
    prs = open_presentation(template)
    remove_all_slides(prs)
    for entry in catalog["layouts"]:
        layout = prs.slide_masters[entry["master"]].slide_layouts[entry["index"]]
        slide = prs.slides.add_slide(layout)
        by_idx = {p["idx"]: p for p in entry["placeholders"]}
        for ph in list(slide.placeholders):
            info = by_idx.get(ph.placeholder_format.idx)
            if info is None:
                continue
            # 枠の範囲が見えるよう破線の矩形を重ねる（画像・グラフ枠は文字が描画されないため）
            outline = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, ph.left, ph.top, ph.width, ph.height)
            outline.fill.background()
            outline.line.color.rgb = RGBColor(0x99, 0x99, 0x99)
            outline.line.dash_style = MSO_LINE.DASH
            outline.shadow.inherit = False
            label = f"[{info['key']}] {entry['id']} {info['type']}"
            if info["type"] in ("PICTURE", "BITMAP", "CHART", "TABLE", "MEDIA_CLIP", "ORG_CHART") or not ph.has_text_frame:
                outline.text_frame.text = label
                outline.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
            else:
                ph.text_frame.text = label
        if not slide.placeholders:
            box = slide.shapes.add_textbox(int(prs.slide_width * 0.05), int(prs.slide_height * 0.45), int(prs.slide_width * 0.9), int(prs.slide_height * 0.1))
            box.text_frame.text = f"{entry['id']} {entry['name']}（プレースホルダー無し）"
    prs.save(target)


def render_thumbs(probe: Path, thumbs_dir: Path, catalog: dict) -> list[str]:
    script = Path(__file__).resolve().parent / "render_pptx.py"
    tmp_out = Path(tempfile.mkdtemp(prefix="decklab-thumbs-"))
    proc = subprocess.run(
        [sys.executable, str(script), str(probe), "--out", str(tmp_out), "--dpi", "60"],
        capture_output=True,
        text=True,
    )
    if proc.returncode != 0:
        print(f"[警告] サムネを作れなかった: {proc.stderr.strip() or proc.stdout.strip()}", file=sys.stderr)
        return []
    thumbs_dir.mkdir(parents=True, exist_ok=True)
    for old in thumbs_dir.glob("*.png"):
        old.unlink()
    pngs = sorted(tmp_out.glob("slide-*.png"))
    written = []
    for entry, png in zip(catalog["layouts"], pngs):
        target = thumbs_dir / f"{entry['id']}.png"
        shutil.move(str(png), target)
        entry["thumb"] = target.name
        written.append(target)
    sys.path.insert(0, str(script.parent))
    from render_pptx import contact_sheet  # noqa: PLC0415

    contact_sheet(written, thumbs_dir / "sheet.png", cols=4, thumb_w=360)
    shutil.rmtree(tmp_out, ignore_errors=True)
    return [str(p) for p in written]


def write_markdown(catalog: dict, issues: list[dict], purposes: dict[str, str], md_path: Path) -> None:
    t = catalog["theme"]
    fonts = t["fonts"]
    lines = [
        "# スライドテンプレ カタログ",
        "",
        "> `register_template.py` が自動生成。**「用途」だけは Claude（または人）が書き込む**。再登録しても用途はレイアウト名で引き継がれる。",
        "",
        f"- テンプレ: `{catalog['template']}`",
        f"- サイズ: {catalog['slide_size']['width_in']} x {catalog['slide_size']['height_in']} in（{catalog['slide_size']['aspect']}）",
        f"- テーマ: {t.get('name') or '-'}",
        f"- 見出しフォント: latin={fonts['major'].get('latin', '-')} / ea={fonts['major'].get('ea', '-')}",
        f"- 本文フォント: latin={fonts['minor'].get('latin', '-')} / ea={fonts['minor'].get('ea', '-')}",
        "- テーマ色: " + ", ".join(f"{k}=#{v}" for k, v in t["colors"].items() if v),
        "- サムネ一覧: `thumbs/sheet.png`",
        "",
        "## レイアウト一覧",
        "",
        "| ID | レイアウト名 | 流し込めるキー |",
        "| --- | --- | --- |",
    ]
    for e in catalog["layouts"]:
        keys = ", ".join(f"`{p['key']}`" for p in e["placeholders"] if not p["footer"]) or "（なし）"
        lines.append(f"| {e['id']} | {e['name']} | {keys} |")
    lines += ["", "## レイアウト詳細", ""]
    for e in catalog["layouts"]:
        lines.append(f"### {e['id']} {e['name']}")
        lines.append("")
        lines.append(f"- **用途**: {purposes.get(e['name'], 'TODO（サムネを見て、どんな主張・内容に使うかを1行で書く）')}")
        if e.get("thumb"):
            lines.append(f"- サムネ: `thumbs/{e['thumb']}`")
        if e["static_texts"]:
            lines.append("- 固定テキスト: " + " / ".join(e["static_texts"]))
        lines.append("")
        content = [p for p in e["placeholders"] if not p["footer"]]
        if content:
            lines.append("| キー | idx | 種類 | 位置 x,y,w,h（%） | 文字 | 目安（全角/行 × 行） |")
            lines.append("| --- | --- | --- | --- | --- | --- |")
            for p in content:
                box = ", ".join(str(v) for v in p["box_pct"])
                cap = p["capacity"]
                lines.append(
                    f"| `{p['key']}` | {p['idx']} | {p['type']} | {box} | {p['font_pt']:g}pt | {cap['chars_per_line']} × {cap['lines']} |"
                )
            lines.append("")
    if catalog["sample_slides"]:
        lines += ["## テンプレ内のサンプルスライド", "", "| # | レイアウト | タイトル | テキストボックス数 |", "| --- | --- | --- | --- |"]
        for s in catalog["sample_slides"]:
            lines.append(f"| {s['no']} | {s['layout']} | {s['title'] or '-'} | {s['free_text_boxes']} |")
        lines.append("")
    lines += ["## 構造チェック（テンプレ側で直すと流し込み精度が上がる）", ""]
    if issues:
        for it in issues:
            lines.append(f"- [{it['level']}] {it['where']}: {it['message']}")
    else:
        lines.append("- 指摘なし")
    lines.append("")
    md_path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("template", help="template.pptx / .potx")
    ap.add_argument("--out", help="出力先（既定: テンプレと同じディレクトリ）")
    ap.add_argument("--no-render", action="store_true", help="サムネを作らない（LibreOffice が無い環境向け）")
    args = ap.parse_args()

    template = Path(args.template)
    if not template.exists():
        sys.exit(f"テンプレが無い: {template}")
    out = Path(args.out) if args.out else template.parent
    out.mkdir(parents=True, exist_ok=True)
    prs = open_presentation(template)
    catalog, issues = analyze(prs, template)

    if not args.no_render:
        probe = Path(tempfile.mkdtemp(prefix="decklab-probe-")) / "probe.pptx"
        build_probe(template, catalog, probe)
        render_thumbs(probe, out / "thumbs", catalog)

    md_path = out / "layouts.md"
    purposes = read_purposes(md_path)
    (out / "layouts.json").write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    write_markdown(catalog, issues, purposes, md_path)

    print(f"layouts: {len(catalog['layouts'])}  sample slides: {len(catalog['sample_slides'])}")
    print(f"written: {out / 'layouts.json'}, {md_path}")
    todo = [e["name"] for e in catalog["layouts"] if e["name"] not in purposes]
    if todo:
        print(f"用途が未記入のレイアウト: {len(todo)} 件（layouts.md の「用途」を埋めること）")
    print("構造チェック:")
    for it in issues or [{"level": "ok", "where": "-", "message": "指摘なし"}]:
        print(f"  [{it['level']}] {it['where']}: {it['message']}")


if __name__ == "__main__":
    main()
