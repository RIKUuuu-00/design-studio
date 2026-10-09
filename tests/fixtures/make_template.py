#!/usr/bin/env python3
"""E2E テスト用の擬似「社内テンプレ」を作る。

python-pptx の既定テンプレをベースに、テーマの日本語フォントを設定し、
見本スライドを 1 枚入れる。実テンプレの代わりにスクリプト一式を通すためだけのもの。

    python make_template.py out/template.pptx [--ea-font IPAGothic]
"""

import argparse
import re
from pathlib import Path

from pptx import Presentation


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("out")
    ap.add_argument("--ea-font", default="IPAGothic")
    ap.add_argument("--latin-font", default="IPAGothic")
    args = ap.parse_args()

    prs = Presentation()
    master = prs.slide_masters[0]
    for rel in master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            part = rel.target_part
            xml = part.blob.decode("utf-8")
            xml = re.sub(r'<a:latin typeface="[^"]*"', f'<a:latin typeface="{args.latin_font}"', xml)
            xml = re.sub(r'<a:ea typeface="[^"]*"', f'<a:ea typeface="{args.ea_font}"', xml)
            xml = xml.replace('name="Office Theme"', 'name="Sample Corp Theme"')
            part._blob = xml.encode("utf-8")
    # 見本スライド（テンプレによくある「使い方」スライド）
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "テンプレの使い方"
    slide.placeholders[1].text = "ここに本文を入力してください"
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.out)
    print(args.out)


if __name__ == "__main__":
    main()
